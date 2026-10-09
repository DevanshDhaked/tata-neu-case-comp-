"""Render every slide to a 2x PNG, then build a pixel-exact PDF (and an image-based PPTX if python-pptx is installed).
Usage (from the deck folder): python3 tools/render.py
"""
import pathlib
from playwright.sync_api import sync_playwright

deck = pathlib.Path(__file__).resolve().parent.parent
out = deck / "renders"
out.mkdir(exist_ok=True)

with sync_playwright() as p:
    b = p.chromium.launch()
    pg = b.new_page(viewport={"width": 1280, "height": 720}, device_scale_factor=2, reduced_motion="reduce")
    pg.goto((deck / "index.html").as_uri())
    pg.add_style_tag(content=".ui{display:none!important}")
    pg.evaluate("document.fonts.ready")
    n = pg.eval_on_selector_all(".slide", "s => s.length")
    pngs = [out / f"slide-{k}.png" for k in range(1, n + 1)]
    for k, f in enumerate(pngs, 1):
        pg.evaluate(f"location.hash = '#{k}'"); pg.wait_for_timeout(700)
        pg.screenshot(path=str(f))

    # PDF built from the renders so it matches the screen pixel for pixel
    html = out / "_pdf.html"
    html.write_text("<style>@page{size:1280px 720px;margin:0}body{margin:0}img{display:block;width:1280px;height:720px;break-after:page}</style>"
                    + "".join(f'<img src="{f.as_uri()}">' for f in pngs))
    pg = b.new_page(); pg.goto(html.as_uri()); pg.wait_for_load_state("networkidle")
    pg.pdf(path=str(deck / "neu-nights-deck.pdf"), width="1280px", height="720px", print_background=True)
    html.unlink()
    b.close()

try:
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation(); prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    for f in pngs:
        prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(str(f), 0, 0, prs.slide_width, prs.slide_height)
    prs.save(deck / "neu-nights-deck.pptx")
except ImportError:
    print("python-pptx not installed: skipped .pptx")
print("ok", n, "slides")
